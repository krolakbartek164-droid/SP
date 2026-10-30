#!/usr/bin/env python3
"""
Synchronizacja logo drużyn: arkusz (zakładka „Audyt drużyn”) -> folder logos/ w repozytorium.

Co robi:
  * czyta z opublikowanego arkusza (ten sam link .xlsx co strona) kolumny I (nazwa drużyny) i K (link do logo),
  * pobiera każde logo z Dysku Google, zmniejsza do miniatury 256 px i zapisuje jako logos/<id>-<odcisk>.webp,
  * zapisuje logos/logos.json (id pliku z Dysku -> nazwa pliku w repozytorium), z którego korzysta strona,
  * nowe logo = nowy plik; zmienione logo (inna zawartość) = nowy plik i stary znika; usunięte z arkusza = znika,
  * jeśli pobranie się nie uda (brak internetu, limit Google), zostawia poprzednią wersję i nic nie psuje.

Uruchomienie lokalne:  python scripts/sync_logos.py            (pobiera arkusz z internetu)
                       python scripts/sync_logos.py plik.xlsx  (z pliku na dysku)
Wymaga: pip install pillow openpyxl
"""
import hashlib, io, json, os, re, sys, time, urllib.request
from concurrent.futures import ThreadPoolExecutor

SHEET_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vRYlbDfQDdAp5vrfh5pinXGLZPab7mZON4_nERFHxVKTpFrrl_P3p995QoEZRi1IG-83aaHU6XZWLMd/pub?output=xlsx")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "logos")
MANIFEST = os.path.join(OUT, "logos.json")
MAX_PX = 256
UA = {"User-Agent": "Mozilla/5.0 (szostki-logo-sync)"}


def http_get(url, tries=3, timeout=40):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
                return r.read()
        except Exception as e:  # noqa
            last = e
            time.sleep(1.5 * (i + 1))
    raise last


def read_links(source):
    """Zwraca {id_z_dysku: nazwa_druzyny} z arkusza Audyt (kolumna I = drużyna, K = link)."""
    import openpyxl
    data = open(source, "rb").read() if os.path.exists(source) else http_get(source)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    name = next((n for n in wb.sheetnames if "audyt" in n.lower()), None)
    if not name:
        raise RuntimeError("Nie znaleziono zakładki 'Audyt' w arkuszu")
    links = {}
    for row in wb[name].iter_rows(min_row=2, values_only=True):
        team = str(row[8] or "").strip() if len(row) > 8 else ""
        link = str(row[10] or "").strip() if len(row) > 10 else ""
        m = re.search(r"/d/([A-Za-z0-9_-]+)", link)
        if team and m:
            links.setdefault(m.group(1), team)
        elif team and link:
            print(f"  ! pomijam (to nie link z Dysku Google): {team} -> {link[:60]}")
    return links


def fetch_original(file_id):
    return http_get(f"https://lh3.googleusercontent.com/d/{file_id}")


def to_thumbnail(raw):
    from PIL import Image
    im = Image.open(io.BytesIO(raw))
    im.seek(0)
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    im = im.convert("RGBA" if has_alpha else "RGB")
    im.thumbnail((MAX_PX, MAX_PX), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "WEBP", quality=90, method=6)
    return buf.getvalue()


def process(file_id, old_name):
    """Zwraca (nazwa_pliku, bajty_lub_None). None = plik już jest aktualny / pobranie się nie udało."""
    try:
        raw = fetch_original(file_id)
        stamp = hashlib.sha256(raw).hexdigest()[:8]
        name = f"{file_id}-{stamp}.webp"
        if name == old_name and os.path.exists(os.path.join(OUT, name)):
            return name, None                      # bez zmian
        return name, to_thumbnail(raw)             # nowe lub zmienione
    except Exception as e:  # noqa
        print(f"  ! {file_id}: {e}")
        return old_name, None                      # zostaje poprzednia wersja


def main():
    source = sys.argv[1] if len(sys.argv) > 1 else SHEET_URL
    os.makedirs(OUT, exist_ok=True)
    old = json.load(open(MANIFEST, encoding="utf-8")) if os.path.exists(MANIFEST) else {}

    links = read_links(source)
    if not links:
        print("Arkusz nie zawiera żadnych linków do logo – niczego nie zmieniam.")
        return
    print(f"Logo w arkuszu: {len(links)}")

    ids = sorted(links)
    with ThreadPoolExecutor(max_workers=4) as ex:
        results = list(ex.map(lambda i: process(i, old.get(i)), ids))

    new, added, changed = {}, [], []
    for fid, (name, blob) in zip(ids, results):
        if not name:
            continue                               # nigdy nie udało się pobrać i nie było wcześniej
        if blob is not None:
            with open(os.path.join(OUT, name), "wb") as f:
                f.write(blob)
            (changed if fid in old else added).append(links[fid])
        new[fid] = name

    # sprzątanie: pliki, których nie ma już w spisie (zmienione lub usunięte logo)
    keep = set(new.values()) | {"logos.json"}
    removed = 0
    for fn in os.listdir(OUT):
        if fn not in keep and fn.endswith(".webp"):
            os.remove(os.path.join(OUT, fn)); removed += 1

    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump(dict(sorted(new.items())), f, ensure_ascii=False, indent=0, sort_keys=True)
        f.write("\n")
    print(f"Nowe: {len(added)} {added[:8]} | zmienione: {len(changed)} {changed[:8]} | usunięte pliki: {removed} | razem: {len(new)}")


if __name__ == "__main__":
    main()
